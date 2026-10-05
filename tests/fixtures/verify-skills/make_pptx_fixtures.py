"""Regenerates tests/fixtures/verify-skills/deck.pptx (python-pptx).

Run: python tests/fixtures/verify-skills/make_pptx_fixtures.py
Needs python-pptx (pip install python-pptx). The committed bytes are the source of truth
for the tests; this records how they were made. Each slide pins one thing:

  Alpha  shapes of every kind interleaved: text, table, text, group (with a nested group),
         chart, picture, text; notes on two lines
  Beta   a body placeholder with two paragraphs, a line break inside a paragraph, and a
         connector and an empty rectangle (shapes with no text)
  Gamma  a purely numeric text run; its notes part is created FIRST, so notesSlide1.xml
         belongs to Gamma, not to the slide that happens to be called slide1.xml

The slides are then reordered to Gamma, Alpha, Beta, so file names (slide3, slide1, slide2)
do not give the presentation order.
"""

import io
import struct
import zlib
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches

here = Path(__file__).parent


def png(width: int, height: int) -> io.BytesIO:
  raw = b"".join(b"\x00" + b"\xff\x80\x00" * width for _ in range(height))

  def chunk(kind: bytes, data: bytes) -> bytes:
    body = kind + data
    return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

  ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
  return io.BytesIO(
      b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")
  )


prs = Presentation()
layouts = {layout.name: layout for layout in prs.slide_layouts}

# --- Alpha
alpha = prs.slides.add_slide(layouts["Title Only"])
alpha.shapes.title.text = "Alpha"
alpha.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4), Inches(0.5)).text_frame.text = "first"
table = alpha.shapes.add_table(2, 2, Inches(0.5), Inches(2.2), Inches(4), Inches(1)).table
for (r, c), text in {(0, 0): "k", (0, 1): "v", (1, 0): "n", (1, 1): "7"}.items():
  table.cell(r, c).text = text
alpha.shapes.add_textbox(Inches(0.5), Inches(3.4), Inches(4), Inches(0.5)).text_frame.text = "second"
group = alpha.shapes.add_group_shape()
group.shapes.add_textbox(Inches(5), Inches(1.5), Inches(3), Inches(0.5)).text_frame.text = "grouped one"
inner = group.shapes.add_group_shape()
inner.shapes.add_textbox(Inches(5), Inches(2.1), Inches(3), Inches(0.5)).text_frame.text = "grouped two"
chart_data = CategoryChartData()
chart_data.categories = ["Q1", "Q2", "Q3"]
chart_data.add_series("Sales", (10, 12.5, 9))
alpha.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(4), Inches(4), Inches(2.5), chart_data
)
alpha.shapes.add_picture(png(8, 4), Inches(5), Inches(4), Inches(2), Inches(1))
alpha.shapes.add_textbox(Inches(0.5), Inches(6.6), Inches(4), Inches(0.5)).text_frame.text = "third"

# --- Beta
beta = prs.slides.add_slide(layouts["Title and Content"])
beta.shapes.title.text = "Beta"
body = beta.placeholders[1].text_frame
body.text = "Bullet 1"
body.add_paragraph().text = "Bullet 2"
broken = beta.shapes.add_textbox(Inches(0.5), Inches(5), Inches(4), Inches(1)).text_frame
broken.text = "line one"
paragraph = broken.paragraphs[0]
paragraph.add_line_break()
paragraph.add_run().text = "line two"
beta.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(5), Inches(5), Inches(7), Inches(5))
beta.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5), Inches(5.5), Inches(1), Inches(0.5))

# --- Gamma (notes created before Alpha's, so notesSlide1.xml is Gamma's)
gamma = prs.slides.add_slide(layouts["Title Only"])
gamma.shapes.title.text = "Gamma"
gamma.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4), Inches(0.5)).text_frame.text = "2024"
gamma.notes_slide.notes_text_frame.text = "Notes for Gamma"
alpha.notes_slide.notes_text_frame.text = "Notes for Alpha\nsecond notes line"

# Presentation order: Gamma, Alpha, Beta
ids = list(prs.slides._sldIdLst)
prs.slides._sldIdLst.remove(ids[2])
prs.slides._sldIdLst.insert(0, ids[2])

prs.save(here / "deck.pptx")
