"""Regenerates tests/fixtures/verify-skills/charts.pptx (python-pptx).

Run: python tests/fixtures/verify-skills/make_pptx_charts_fixture.py
Needs python-pptx (pip install python-pptx). The committed bytes are the source of truth for
the tests; this records how they were made. It pins the parts of pptx_text.ts that must tell
cases apart the way python-pptx does:

  Charts  one chart per kind whose type name differs (stacked or not, markers or not,
          exploded or not, 3-D bubbles, ...), with categories that need quoting in a Python
          list, a missing value and floats that print in exponent form
  Shapes  a slide with no title: padded and empty paragraphs, a line break, pictures whose
          size is an exact tie at one decimal (0.25 in and 1.25 in), a table with a merged
          cell, and notes that hold only spaces
"""

import io
import struct
import zlib
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE as X
from pptx.util import Inches

here = Path(__file__).parent


def png(width: int, height: int) -> io.BytesIO:
  raw = b"".join(b"\x00" + b"\x10\x20\x30" * width for _ in range(height))

  def chunk(kind: bytes, data: bytes) -> bytes:
    body = kind + data
    return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

  ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
  return io.BytesIO(
      b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")
  )


prs = Presentation()
layouts = {layout.name: layout for layout in prs.slide_layouts}

charts = prs.slides.add_slide(layouts["Title Only"])
charts.shapes.title.text = "Charts"
CATEGORIES = ["Q1", "it's", 'say "hi"', "back" + chr(92) + "slash", "caf" + chr(0xE9)]
VALUES = (1, 2.5, None, 1e-05, 1e16)
for kind in (
    X.COLUMN_STACKED,
    X.BAR_CLUSTERED,
    X.LINE,
    X.LINE_MARKERS,
    X.PIE_EXPLODED,
    X.DOUGHNUT,
    X.AREA_STACKED,
    X.RADAR_FILLED,
):
  data = CategoryChartData()
  data.categories = CATEGORIES
  data.add_series("First", VALUES)
  data.add_series("Second " + chr(0x2014) + " two", (3, 4, 5, 6, 7))
  charts.shapes.add_chart(kind, Inches(0.5), Inches(1.5), Inches(3), Inches(2), data)

scatter = XyChartData()
points = scatter.add_series("points")
for x, y in ((1, 2), (2, 3.5), (3, 1)):
  points.add_data_point(x, y)
charts.shapes.add_chart(
    X.XY_SCATTER_LINES_NO_MARKERS, Inches(0.5), Inches(1.5), Inches(3), Inches(2), scatter
)

bubbles = BubbleChartData()
blobs = bubbles.add_series("blobs")
blobs.add_data_point(1, 2, 3)
blobs.add_data_point(2, 3, 4)
charts.shapes.add_chart(
    X.BUBBLE_THREE_D_EFFECT, Inches(0.5), Inches(1.5), Inches(3), Inches(2), bubbles
)

many = CategoryChartData()
many.categories = [f"c{i}" for i in range(15)]
many.add_series("long", list(range(15)))
charts.shapes.add_chart(X.LINE, Inches(0.5), Inches(1.5), Inches(3), Inches(2), many)

shapes = prs.slides.add_slide(layouts["Blank"])
frame = shapes.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text_frame
frame.text = "  padded  "
frame.add_paragraph().text = ""
broken = frame.add_paragraph()
broken.text = "a" + chr(0xA0) + "b" + chr(0x2028) + "c"
broken.add_line_break()
broken.add_run().text = "after break"
frame.add_paragraph().text = "tail  "
shapes.shapes.add_picture(png(4, 4), 0, 0, Inches(0.25), Inches(0.75))
shapes.shapes.add_picture(png(4, 4), 0, 0, Inches(1.25), Inches(2.5))
table = shapes.shapes.add_table(3, 3, 0, 0, Inches(4), Inches(2)).table
table.cell(0, 0).merge(table.cell(0, 1))
table.cell(0, 0).text = "merged\nhead"
table.cell(1, 0).text = "  x "
table.cell(2, 2).text = "tail"
shapes.notes_slide.notes_text_frame.text = "   "

prs.save(here / "charts.pptx")
