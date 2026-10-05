"""Regenerates tests/fixtures/verify-skills/view.docx (python-docx).

Run: python tests/fixtures/verify-skills/make_docx_view_fixture.py
Needs python-docx (pip install python-docx). The committed bytes are the source of truth for
the tests; this records how they were made. Each block pins one rule of the .docx text view:

  headings     styles are told apart by NAME ("Heading 1", "Title"), not by style id
  runs         a purely numeric run, spaces between runs, `&` and `<`, a tab, a line break, a
               page break (which leaves no text) and a hyperlink all read as python-docx reads them
  table        a column span and a row span repeat their text in every grid cell they cover
"""

from pathlib import Path

import docx
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

here = Path(__file__).parent
doc = docx.Document()

doc.add_heading("Report", level=1)
doc.add_paragraph("Quarterly", style="Title")

totals = doc.add_paragraph()
totals.add_run("Total ")
totals.add_run("2024")
totals.add_run(" units")

greeting = doc.add_paragraph()
greeting.add_run("Hello ")
greeting.add_run("world")
doc.add_paragraph("Fish & chips <b>")

breaks = doc.add_paragraph("a\tb")
breaks.add_run().add_break()
breaks.add_run("c")
breaks.add_run().add_break(WD_BREAK.PAGE)
breaks.add_run("d")

linked = doc.add_paragraph("See ")
linked._p.append(
    parse_xml(
        f'<w:hyperlink {nsdecls("w")} w:anchor="top"><w:r><w:t>the link</w:t></w:r></w:hyperlink>'
    )
)
linked.add_run(" now")

doc.add_paragraph("")  # an empty paragraph leaves no line

# Name and id disagree: the name decides.
by_id = doc.styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
by_id.element.set(qn("w:styleId"), "HeadingLike")
doc.add_paragraph("heading-looking id, plain name", style=by_id)
by_name = doc.styles.add_style("Title Alt", WD_STYLE_TYPE.PARAGRAPH)
by_name.element.set(qn("w:styleId"), "Plain7")
doc.add_paragraph("plain-looking id, title name", style=by_name)

table = doc.add_table(rows=3, cols=3)
table.cell(0, 0).merge(table.cell(0, 1))  # column span
table.cell(1, 0).merge(table.cell(2, 0))  # row span
table.cell(0, 0).text = "head"
table.cell(0, 2).text = "2024"
table.cell(1, 0).text = "tall"
table.cell(1, 1).text = "a\nb"
table.cell(1, 2).text = "  x  "
table.cell(2, 1).text = "y"
table.cell(2, 2).text = ""

doc.save(here / "view.docx")
